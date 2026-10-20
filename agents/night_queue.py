"""야간 작업 큐: Linear에서 '야간작업' 라벨이 붙은 이슈를 Claude로 처리하고 결과를 댓글로 남긴다.

실행: python -m agents.night_queue               (라벨 붙은 이슈 전체 처리, GitHub Actions용)
      python -m agents.night_queue --dry-run     (처리할 이슈 목록만 확인, 무료)
      python -m agents.night_queue --issue RES-12 (라벨과 관계없이 이 이슈만 처리)

저장소가 공개라서, 이 스크립트는 작업 내용·제목을 로그나 저장소에 남기지 않는다 (이슈 번호만 출력).
"""
import argparse
import base64
import html
import io
import os
import re
import sys
import zipfile

import requests

from agents import config, linear
from agents.common import Meter, log, notify_error, research

SYSTEM = f"""너는 로봇·AI 분야 박사과정 연구자의 연구 보조원이다. 사용자가 Linear 이슈로 맡긴 일을 밤사이 처리한다.
사용자는 아침에 결과를 읽고 바로 쓸 수 있어야 한다. 한국어로 쓰고, 기술 용어는 원래 표기를 유지한다.

사용자의 주요 연구 (관련 있을 때만 참고):
{config.RESEARCH_CONTEXT}

작업 원칙:
- 이슈 제목과 설명이 지시사항이다. 이전 댓글에 수정 요청이 있으면 그것을 우선 반영한다.
- 지시가 모호하면 가장 그럴듯한 해석으로 끝까지 해내고, 어떻게 해석했는지 '확인이 필요한 점'에 적는다.
  질문만 남기고 멈추지 않는다.
- 문서 수정 요청이면 바꿀 부분을 '수정 전 / 수정 후' 형식으로 보여주거나, 바로 붙여넣을 수 있는 완성 텍스트로 준다.
- 두 문서 대조 요청이면 수치, 용어, 누락 내용의 불일치를 위치와 함께 표로 정리한다.
- 첨부 파일에 없는 수치나 사실을 지어내지 않는다. 모르면 모른다고 쓴다.
- 논문·웹 자료를 근거로 들 때는 출처 URL과 함께 원문 문구를 한 글자도 바꾸지 말고 따옴표로 옮긴다 (15단어 이내).
  사용자가 원문에서 그 문구를 검색해 확인한다. 정확히 옮길 수 없으면 인용하지 말고 출처만 적는다.
- 웹 검색은 최신 정보나 문헌 확인이 꼭 필요할 때만 쓴다.

출력 형식 (Linear 댓글 마크다운):
## 요약
(결과를 3줄 이내로)

## 결과
(본문)

## 확인이 필요한 점
(가정한 것, 사용자가 직접 확인할 것. 없으면 '없음')"""

TEXT_EXT = {".txt", ".md", ".csv", ".tsv", ".json", ".py", ".m", ".yaml", ".yml", ".tex", ".bib", ".log"}
IMG_TYPES = {".png": "image/png", ".jpg": "image/jpeg", ".jpeg": "image/jpeg", ".gif": "image/gif", ".webp": "image/webp"}
MAX_TEXT_CHARS = 300_000


# ---------------------------------------------------------------- 첨부 파일 → Claude 입력
def _ext(name, ctype, url):
    for s in (name, url):
        m = re.search(r"(\.[A-Za-z0-9]{1,5})$", s.split("?")[0])
        if m:
            return m.group(1).lower()
    guess = {"application/pdf": ".pdf", "image/png": ".png", "image/jpeg": ".jpg", "image/gif": ".gif",
             "image/webp": ".webp", "text/csv": ".csv", "text/plain": ".txt",
             "application/vnd.openxmlformats-officedocument.wordprocessingml.document": ".docx",
             "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": ".xlsx"}
    return guess.get(ctype.split(";")[0].strip(), "")


def _decode(data):
    for enc in ("utf-8", "cp949"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            pass
    return data.decode("utf-8", errors="replace")


def _docx_text(data):
    import docx
    d = docx.Document(io.BytesIO(data))
    out = [p.text for p in d.paragraphs if p.text.strip()]
    for ti, t in enumerate(d.tables, 1):
        out.append(f"\n[표 {ti}]")
        out += [" | ".join(c.text.strip() for c in row.cells) for row in t.rows]
    return "\n".join(out)


def _hwpx_text(data):
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        parts = sorted(n for n in z.namelist() if re.match(r"Contents/section\d+\.xml$", n))
        out = []
        for n in parts:
            xml = z.read(n).decode("utf-8", errors="replace")
            xml = re.sub(r"</hp:p>", "\n", xml)
            out.append(html.unescape(re.sub(r"<[^>]+>", "", xml)))
    return "\n".join(out).strip()


def _xlsx_text(data):
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    out = []
    for ws in wb.worksheets:
        out.append(f"[시트: {ws.title}]")
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i >= 2000:
                out.append("... (2000행 이후 생략)")
                break
            if any(v is not None for v in row):
                out.append("\t".join("" if v is None else str(v) for v in row))
    return "\n".join(out)


def file_to_blocks(name, url):
    """첨부 1개를 Claude 메시지 블록으로 바꾼다. (blocks, 메모)"""
    try:
        data, ctype = linear.download(url, config.NIGHT_MAX_FILE_MB)
    except Exception as e:
        return [], f"{name or '첨부'}: 내려받기 실패 ({e})"
    ext = _ext(name, ctype, url)
    label = name or f"첨부{ext}"
    try:
        if ext == ".pdf":
            return [{"type": "document", "title": label[:200],
                     "source": {"type": "base64", "media_type": "application/pdf",
                                "data": base64.b64encode(data).decode()}}], ""
        if ext in IMG_TYPES:
            return [{"type": "image", "source": {"type": "base64", "media_type": IMG_TYPES[ext],
                                                 "data": base64.b64encode(data).decode()}}], ""
        if ext == ".docx":
            text = _docx_text(data)
        elif ext == ".hwpx":
            text = _hwpx_text(data)
        elif ext in (".xlsx", ".xlsm"):
            text = _xlsx_text(data)
        elif ext in TEXT_EXT:
            text = _decode(data)
        elif ext in (".hwp", ".doc", ".ppt", ".pptx", ".xls"):
            return [], f"{label}: 이 형식은 읽을 수 없습니다. HWPX, DOCX, PDF, XLSX로 저장해 다시 첨부하세요"
        else:
            return [], f"{label}: 지원하지 않는 형식입니다"
    except Exception as e:
        return [], f"{label}: 내용 추출 실패 ({e})"
    if len(text) > MAX_TEXT_CHARS:
        text = text[:MAX_TEXT_CHARS] + "\n... (길이 제한으로 이후 생략)"
    return [{"type": "text", "text": f"<첨부 파일 이름=\"{label}\">\n{text}\n</첨부 파일>"}], ""


def build_content(issue):
    comments = sorted(issue["comments"]["nodes"], key=lambda c: c["createdAt"])
    texts = [issue.get("description") or ""] + [c["body"] for c in comments]
    blocks, notes = [], []
    for name, url in linear.upload_links("\n".join(texts)):
        b, note = file_to_blocks(name, url)
        blocks += b
        if note:
            notes.append(note)
    history = "\n\n".join(f"[{c['createdAt'][:16].replace('T', ' ')} {(c.get('user') or {}).get('name', '')}]\n{c['body']}"
                          for c in comments)[-60_000:]
    task = (f"<이슈 번호>{issue['identifier']}</이슈 번호>\n<제목>{issue['title']}</제목>\n"
            f"<설명>\n{issue.get('description') or '(설명 없음)'}\n</설명>\n"
            + (f"<이전 댓글>\n{history}\n</이전 댓글>\n" if history else "")
            + (f"<첨부 처리 메모>\n" + "\n".join(notes) + "\n</첨부 처리 메모>\n" if notes else "")
            + "\n위 이슈의 작업을 수행하라.")
    return blocks + [{"type": "text", "text": task}], notes


# ---------------------------------------------------------------- 처리
def split_comment(body, limit=60_000):
    if len(body) <= limit:
        return [body]
    parts, cur = [], ""
    for para in body.split("\n\n"):
        if len(cur) + len(para) + 2 > limit and cur:
            parts.append(cur)
            cur = ""
        cur += para + "\n\n"
    parts.append(cur)
    return [f"({i}/{len(parts)})\n\n{p}" for i, p in enumerate(parts, 1)]


def finish_issue(issue, team_obj, done_label_id, success):
    labels = [l["id"] for l in issue["labels"]["nodes"] if l["name"] != config.NIGHT_LABEL]
    if success and done_label_id and done_label_id not in labels:
        labels.append(done_label_id)
    fields = {"labelIds": labels}
    sid = state_id(team_obj, config.LINEAR_REVIEW_STATE) if success else None
    if sid:
        fields["stateId"] = sid
    linear.update_issue(issue["id"], **fields)


def state_id(team_obj, name):
    return linear.state_id(team_obj, name) if team_obj else None


def process(issue, team_obj, done_label_id):
    meter = Meter()
    try:
        content, notes = build_content(issue)
        result = research(SYSTEM, content, config.NIGHT_MAX_SEARCHES, meter).strip()
        if not result:
            raise RuntimeError("Claude가 빈 결과를 돌려줬습니다")
        head = f"🌙 **야간 작업 결과** (Claude {config.MODEL}, 비용 약 ${meter.cost:.2f})\n\n"
        for part in split_comment(head + result):
            linear.comment(issue["id"], part)
        finish_issue(issue, team_obj, done_label_id, True)
        summary = _summary(result)
        return {"ok": True, "cost": meter.cost, "summary": summary, "notes": notes}
    except Exception as e:
        try:
            linear.comment(issue["id"], f"⚠️ **야간 작업 실패**: {e}\n\n고친 뒤 `{config.NIGHT_LABEL}` 라벨을 다시 붙이면 다음 실행 때 재시도합니다.")
            finish_issue(issue, team_obj, done_label_id, False)
        except Exception:
            pass
        return {"ok": False, "cost": meter.cost, "summary": str(e)[:300], "notes": []}


def _summary(result):
    m = re.search(r"##\s*요약\s*\n(.*?)(?:\n##\s|\Z)", result, re.S)
    text = (m.group(1) if m else result).strip()
    return text[:500]


# ---------------------------------------------------------------- Slack
def slack_report(results):
    webhook = os.getenv("SLACK_WEBHOOK_URL_NIGHT") or os.getenv("SLACK_WEBHOOK_URL")
    if not webhook or not results:
        return
    ok = sum(1 for _, r in results if r["ok"])
    total = sum(r["cost"] for _, r in results)
    blocks = [{"type": "header", "text": {"type": "plain_text", "text": f"🌙 야간 작업 {ok}/{len(results)}건 완료"}},
              {"type": "divider"}]
    for issue, r in results:
        mark = "✅" if r["ok"] else "⚠️ 실패"
        title = issue["title"].replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
        summ = r["summary"].replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
        txt = f"{mark} *<{issue['url']}|{issue['identifier']} {title}>*\n{summ}"
        if r["notes"]:
            txt += "\n_첨부 메모: " + "; ".join(r["notes"])[:400] + "_"
        blocks.append({"type": "section", "text": {"type": "mrkdwn", "text": txt[:2900]}})
    blocks += [{"type": "divider"},
               {"type": "context", "elements": [{"type": "mrkdwn",
                "text": f"전체 결과는 Linear 댓글에 있습니다   |   비용 약 ${total:.2f}"}]}]
    requests.post(webhook, json={"text": f"야간 작업 {ok}/{len(results)}건 완료", "blocks": blocks[:50],
                                 "unfurl_links": False}, timeout=30).raise_for_status()


# ---------------------------------------------------------------- 실행
def main():
    ap = argparse.ArgumentParser(description="Linear 야간 작업 큐")
    ap.add_argument("--dry-run", action="store_true", help="처리할 이슈 번호만 출력 (무료)")
    ap.add_argument("--issue", help="라벨과 관계없이 이 이슈만 처리 (예: RES-12)")
    args = ap.parse_args()

    if not os.getenv("LINEAR_API_KEY", "").startswith("lin_api_") or not config.LINEAR_TEAM_KEY:
        sys.exit("LINEAR_API_KEY 또는 LINEAR_TEAM_KEY가 비어 있습니다. python -m agents.linear_setup을 먼저 실행하세요")

    try:
        issues = [linear.issue_by_identifier(args.issue)] if args.issue else linear.queued_issues(config.NIGHT_MAX_ISSUES)
        issues = [i for i in issues if i]
        log(f"[info] 처리 대상 {len(issues)}건: {', '.join(i['identifier'] for i in issues) or '없음'}")
        if args.dry_run or not issues:
            return
        if not os.getenv("ANTHROPIC_API_KEY", "").startswith("sk-ant-"):
            raise RuntimeError("ANTHROPIC_API_KEY가 없거나 형식이 다릅니다")
        team_obj = linear.team()
        done_label_id, _ = linear.ensure_label(config.NIGHT_DONE_LABEL, team_obj["id"])
        results = []
        for issue in issues:
            r = process(issue, team_obj, done_label_id)
            log(f"[{'ok' if r['ok'] else 'fail'}] {issue['identifier']} 비용 약 ${r['cost']:.3f}")
            results.append((issue, r))
        slack_report(results)
    except Exception as e:
        notify_error(os.getenv("SLACK_WEBHOOK_URL_NIGHT") or os.getenv("SLACK_WEBHOOK_URL", ""), "야간 작업 큐", e)
        raise


if __name__ == "__main__":
    main()
