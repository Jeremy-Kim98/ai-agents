"""네트워크 없이 Linear 연동 흐름을 검증한다 (가짜 Linear, 가짜 Claude, 가짜 Slack).

실행: python -m tests.test_offline
"""
import io
import json
import os
import sys
import zipfile

os.environ.update({"LINEAR_API_KEY": "lin_api_test", "LINEAR_TEAM_KEY": "RES",
                   "ANTHROPIC_API_KEY": "sk-ant-test", "SLACK_WEBHOOK_URL": "https://hooks.slack.com/services/x"})

from agents import config, linear, night_queue  # noqa: E402

config.LINEAR_TEAM_KEY = "RES"
PASS = []


def check(cond, msg):
    if not cond:
        print(f"✗ {msg}")
        sys.exit(1)
    PASS.append(msg)
    print(f"✓ {msg}")


# ---------------------------------------------------------------- 가짜 파일
def make_docx():
    import docx
    d = docx.Document()
    d.add_paragraph("분사 작동 높이는 80cm로 한다.")
    t = d.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text, t.rows[0].cells[1].text = "노즐", "TP00009"
    b = io.BytesIO(); d.save(b); return b.getvalue()


def make_hwpx():
    b = io.BytesIO()
    with zipfile.ZipFile(b, "w") as z:
        z.writestr("Contents/section0.xml",
                   '<hs:sec><hp:p><hp:run><hp:t>분사 높이 40cm</hp:t></hp:run></hp:p>'
                   '<hp:p><hp:run><hp:t>판단·연산부 &amp; CCD</hp:t></hp:run></hp:p></hs:sec>')
    return b.getvalue()


def make_xlsx():
    import openpyxl
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "측정"
    ws.append(["노즐", "VMD"]); ws.append(["TP00009", 152.3])
    b = io.BytesIO(); wb.save(b); return b.getvalue()


FILES = {
    "https://uploads.linear.app/a/b/plan": (make_hwpx(), "application/octet-stream"),
    "https://uploads.linear.app/a/b/report": (make_docx(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
    "https://uploads.linear.app/a/b/data": (make_xlsx(), "application/octet-stream"),
    "https://uploads.linear.app/a/b/fig": (b"\x89PNG fake", "image/png"),
    "https://uploads.linear.app/a/b/paper": (b"%PDF-1.4 fake", "application/pdf"),
    "https://uploads.linear.app/a/b/old": (b"HWP old", "application/octet-stream"),
}
linear.download = lambda url, mb: FILES[url]

# ---------------------------------------------------------------- 1. 첨부 추출
desc = ("중간보고서와 계획서를 대조해줘.\n[계획서.hwpx](https://uploads.linear.app/a/b/plan)\n"
        "[중간보고서.docx](https://uploads.linear.app/a/b/report)\n[측정.xlsx](https://uploads.linear.app/a/b/data)\n"
        "![](https://uploads.linear.app/a/b/fig)\n[논문.pdf](https://uploads.linear.app/a/b/paper)\n"
        "[옛날.hwp](https://uploads.linear.app/a/b/old)\nhttps://uploads.linear.app/a/b/plan")
links = linear.upload_links(desc)
check(len(links) == 6, "첨부 링크 6개 추출 (중복 제거)")
blocks = {n: night_queue.file_to_blocks(n, u) for n, u in links}
check("40cm" in blocks["계획서.hwpx"][0][0]["text"] and "&" in blocks["계획서.hwpx"][0][0]["text"], "HWPX 본문 추출")
check("80cm" in blocks["중간보고서.docx"][0][0]["text"] and "TP00009" in blocks["중간보고서.docx"][0][0]["text"], "DOCX 본문·표 추출")
check("152.3" in blocks["측정.xlsx"][0][0]["text"], "XLSX 시트 추출")
check(blocks[""][0][0]["type"] == "image", "이미지는 image 블록")
check(blocks["논문.pdf"][0][0]["type"] == "document", "PDF는 document 블록")
check(blocks["옛날.hwp"][0] == [] and "HWPX" in blocks["옛날.hwp"][1], "HWP는 안내 메모와 함께 건너뜀")

# ---------------------------------------------------------------- 2. 가짜 Linear
STATE = {"comments": [], "updates": [], "created": [], "labels": {"야간완료": "lab-done", "레이더": None}}
ISSUE = {"id": "iss-1", "identifier": "RES-12", "title": "보고서 대조", "description": desc,
         "url": "https://linear.app/x/issue/RES-12", "team": {"id": "team-1", "key": "RES"},
         "state": {"name": "Todo", "type": "unstarted"},
         "labels": {"nodes": [{"id": "lab-night", "name": "야간작업"}, {"id": "lab-x", "name": "보고서"}]},
         "comments": {"nodes": [{"body": "표로 정리해줘", "createdAt": "2026-10-19T10:00:00.000Z", "user": {"name": "나"}}]}}
TEAM = {"id": "team-1", "key": "RES", "name": "연구",
        "states": {"nodes": [{"id": "st-rev", "name": "In Review", "type": "started"},
                             {"id": "st-todo", "name": "Todo", "type": "unstarted"}]}}


def fake_gql(q, v=None):
    v = v or {}
    if "teams(filter" in q:
        return {"teams": {"nodes": [TEAM]}}
    if "issueLabels" in q:
        lid = STATE["labels"].get(v["name"])
        return {"issueLabels": {"nodes": [{"id": lid, "name": v["name"], "team": {"id": "team-1"}}] if lid else []}}
    if "issueLabelCreate" in q:
        STATE["labels"][v["input"]["name"]] = "lab-new"
        return {"issueLabelCreate": {"success": True, "issueLabel": {"id": "lab-new"}}}
    if "issues(first: $n" in q:
        assert v["label"] == "야간작업" and v["team"] == "RES"
        return {"issues": {"nodes": [ISSUE]}}
    if "commentCreate" in q:
        STATE["comments"].append(v["input"]); return {"commentCreate": {"success": True}}
    if "issueUpdate" in q:
        STATE["updates"].append(v); return {"issueUpdate": {"success": True}}
    if "issueCreate" in q:
        STATE["created"].append(v["input"])
        n = len(STATE["created"])
        return {"issueCreate": {"success": True, "issue": {"id": f"new-{n}", "identifier": f"RES-{100 + n}",
                                                            "url": f"https://linear.app/x/issue/RES-{100 + n}"}}}
    if "updatedAt" in q:
        return {"issues": {"nodes": [{"title": "[레이더] YOLO 새 버전", "state": {"type": "completed"}},
                                     {"title": "[레이더] 휴머노이드 발표", "state": {"type": "canceled"}}]}}
    raise AssertionError("예상 못 한 쿼리: " + q[:80])


linear.gql = fake_gql

# ---------------------------------------------------------------- 3. 가짜 Claude, 가짜 Slack
SEEN = {}


def fake_research(system, content, max_searches, meter):
    SEEN["content"] = content
    meter.cost += 0.42
    return "## 요약\n분사 높이가 계획서 40cm, 보고서 80cm로 다릅니다.\n\n## 결과\n| 항목 | 계획서 | 보고서 |\n\n## 확인이 필요한 점\n없음"


night_queue.research = fake_research
SLACK = []
night_queue.requests.post = lambda url, json=None, timeout=None: (SLACK.append(json), type("R", (), {"raise_for_status": lambda s: None})())[1]

sys.argv = ["night_queue"]
night_queue.main()

task_text = SEEN["content"][-1]["text"]
check("RES-12" in task_text and "표로 정리해줘" in task_text and "HWPX, DOCX" in task_text, "지시·이전 댓글·첨부 메모가 Claude 입력에 포함")
check(sum(1 for b in SEEN["content"] if b["type"] in ("document", "image")) == 2, "PDF·이미지 블록 전달")
check(len(STATE["comments"]) == 1 and "야간 작업 결과" in STATE["comments"][0]["body"], "결과를 Linear 댓글로 작성")
upd = STATE["updates"][-1]["input"]
check("lab-night" not in upd["labelIds"] and "lab-done" in upd["labelIds"] and "lab-x" in upd["labelIds"],
      "야간작업 라벨 제거, 야간완료 추가, 기존 라벨 유지")
check(upd.get("stateId") == "st-rev", "상태를 In Review로 이동")
check(SLACK and "1/1" in SLACK[-1]["text"] and "40cm" in json.dumps(SLACK[-1], ensure_ascii=False), "Slack 요약 전송")

# 실패 경로
STATE["comments"].clear(); STATE["updates"].clear(); SLACK.clear()
night_queue.research = lambda *a, **k: (_ for _ in ()).throw(RuntimeError("크레딧 부족"))
night_queue.main()
check("야간 작업 실패" in STATE["comments"][-1]["body"], "실패 시 Linear에 실패 댓글")
upd = STATE["updates"][-1]["input"]
check("lab-night" not in upd["labelIds"] and "lab-done" not in upd["labelIds"] and "stateId" not in upd,
      "실패 시 라벨만 제거, 상태 유지 (재시도는 라벨 다시 붙이기)")
check("0/1" in SLACK[-1]["text"], "Slack에 실패 표시")

# 긴 결과 분할
parts = night_queue.split_comment("가" * 50_000 + "\n\n" + "나" * 50_000)
check(len(parts) == 2 and parts[0].startswith("(1/2)"), "긴 결과는 댓글 여러 개로 분할")

# ---------------------------------------------------------------- 4. 레이더 연동
items = [{"title": "새 검출 모델", "summary": "요약", "why": "이유", "apply": "검출 단계", "category": "검출",
          "source": "GitHub", "date": "2026-10-19", "url": "https://github.com/x", "quote": "faster",
          "importance": 3, "quote_status": "ok"},
         {"title": "덜 중요", "summary": "s", "why": "", "apply": "", "category": "", "source": "", "date": "",
          "url": "https://y", "quote": "", "importance": 2}]
linear.radar_to_linear(items)
check(len(STATE["created"]) == 1 and STATE["created"][0]["title"].startswith("[레이더]"), "🔥 항목만 Linear 이슈 생성")
check(items[0].get("linear_id") == "RES-101", "레이더 항목에 이슈 번호 연결")
done, dropped = linear.radar_feedback()
check(done == ["YOLO 새 버전"] and dropped == ["휴머노이드 발표"], "Linear 진행/취소 기록을 피드백으로 변환")

from agents import robot_radar  # noqa: E402
fb = robot_radar.feedback_text()
check("관심 있어 진행한 주제" in fb and "휴머노이드" in fb, "레이더 프롬프트에 피드백 반영")

print(f"\n전부 통과 ({len(PASS)}개)")
